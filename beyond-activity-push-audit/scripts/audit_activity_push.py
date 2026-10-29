#!/usr/bin/env python3
"""Beyond 活动 Push 气泡 / Tab 配置核对。

只读扫描 ``000_活动push_气泡推送表.csv`` 与 ``Data/TableCfg/JSON/<ver>d99*``，
输出 md / json（装了 openpyxl 时再出 xlsx）。不会改表。

字段名、pushType 取值、三服时区都在 ``DEFAULT_SCHEMA``；工程实际字段不同时，
在本脚本同目录放 ``push_audit_schema.json``（或 ``--schema``）覆盖，无需改代码。
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

CSV_NAME = "000_活动push_气泡推送表.csv"
JSON_REL = Path("Data") / "TableCfg" / "JSON"
CSV_CACHE_REL = Path("Shared") / "Data" / "Tables" / "~CsvCache"
REPORT_REL = Path("Temp") / "ForAIReports"
PYTHON_TEST_ROOT = Path(os.environ.get("BEYOND_PYTHON_TEST_ROOT", r"D:\python-test"))
PYTHON_TEST_CONFIG = "rule_project_config.json"
SCHEMA_FILE = "push_audit_schema.json"
CURRENT_RE = re.compile(r"^CURRENT(d\d+)?$", re.IGNORECASE)
JSON_DIR_RE = re.compile(r"^(v\d+d\d+)d99", re.IGNORECASE)
UTC = dt.timezone.utc

DEFAULT_SCHEMA: dict[str, Any] = {
    "csv_fields": {
        "id": ["id", "pushId", "bubbleId"],
        "activity": ["activityId", "actId", "activity"],
        "push_type": ["pushType"],
        "carrier": ["pushCategory", "carrier", "showType"],
        "sort": ["bubbleSortId", "sortId"],
        "offset": ["offset", "timeOffset", "activateOffset", "startOffset"],
        "end_time": ["activityEndTime"],
        "version_start": ["versionStart"],
        "version_end": ["versionEnd"],
    },
    # pushType 原值 -> bubble_open / bubble_update / bubble_end / tab_open / tab_update / tab_end。
    # 未命中时按关键字猜（tab/页签、bubble/气泡、update/更新、end/临期、open/开放）。
    "push_type_map": {},
    "tables": {
        "activity": ["ActivityTable"],
        "time_range": ["TimeRangeTable"],
        "cms": ["ActivityCmsTable", "CmsTable", "ActivityCMSTable"],
        "game_entrance": ["GameEntranceTable"],
    },
    "activity_fields": {
        "id": ["id", "activityId"],
        "name": ["name", "activityName", "title"],
        "time_id": ["timeId", "timeRangeId"],
        "open": ["openTime", "startTime"],
        "close": ["closeTime", "endTime"],
        "version_start": ["versionStart"],
        "version_end": ["versionEnd"],
    },
    "time_range_fields": {
        "id": ["id", "timeId"],
        "start": ["startTime", "openTime", "beginTime"],
        "end": ["endTime", "closeTime"],
    },
    "stage_fields": {
        "id": ["id", "entranceId", "cmsId"],
        "activity": ["activityId", "actId"],
        "time_id": ["timeId"],
        "time_offset": ["timeOffset"],
        "close_offset": ["closeOffset", "endOffset"],
    },
    # suffixes：三服专属字段后缀（如 startTimeApac）；找不到时回退到无后缀字段，
    # 无时区的时间字符串按该服 utc_offset_hours 解释（墙钟对齐）。
    "servers": {
        "CN": {"label": "国服", "utc_offset_hours": 8, "suffixes": ["Cn", "CN", "_cn"]},
        "APAC": {"label": "亚太", "utc_offset_hours": 8, "suffixes": ["Apac", "APAC", "_apac", "Asia"]},
        "NAEU": {"label": "欧美", "utc_offset_hours": -5, "suffixes": ["Global", "NaEu", "Na", "Eu", "_global", "Os"]},
    },
    "offset_unit_seconds": 1,
    "end_lead_hours": 72,
}

KIND_LABEL = {
    "bubble_open": "已开放气泡",
    "bubble_update": "更新气泡",
    "bubble_end": "临期气泡",
    "tab_open": "已开放Tab",
    "tab_update": "Update Tab",
    "tab_end": "End Tab",
}
META_TOKENS = {"int", "long", "string", "str", "bool", "float", "double", "list", "array", "enum", "key", "id"}


# --------------------------------------------------------------------------- 基础工具


def merge_schema(base: dict, override: dict) -> dict:
    out = dict(base)
    for key, value in override.items():
        if isinstance(value, dict) and isinstance(out.get(key), dict):
            out[key] = merge_schema(out[key], value)
        else:
            out[key] = value
    return out


def pick(rec: dict, aliases: Iterable[str]) -> Any:
    lower = {str(k).lower(): k for k in rec}
    for alias in aliases:
        key = lower.get(alias.lower())
        if key is not None and rec[key] not in (None, ""):
            return rec[key]
    return None


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def parse_time(value: Any, tz: dt.tzinfo) -> dt.datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return None
    raw = str(value).strip()
    if re.fullmatch(r"-?\d+(\.\d+)?", raw):
        num = float(raw)
        if num > 1e11:
            num /= 1000.0
        if num < 1e8:
            return None
        return dt.datetime.fromtimestamp(num, UTC)
    cleaned = raw.replace("/", "-").replace("Z", "+00:00")
    parsed: dt.datetime | None = None
    try:
        parsed = dt.datetime.fromisoformat(cleaned)
    except ValueError:
        for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d"):
            try:
                parsed = dt.datetime.strptime(cleaned, fmt)
                break
            except ValueError:
                continue
    if parsed is None:
        raise ValueError(f"无法解析时间: {raw}")
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=tz)
    return parsed.astimezone(UTC)


DURATION_RE = re.compile(
    r"^(?P<sign>-)?\s*(?:(?P<d>\d+(?:\.\d+)?)\s*d)?\s*(?:(?P<h>\d+(?:\.\d+)?)\s*h)?\s*"
    r"(?:(?P<m>\d+(?:\.\d+)?)\s*m(?:in)?)?\s*(?:(?P<s>\d+(?:\.\d+)?)\s*s)?$",
    re.IGNORECASE,
)


def parse_duration(value: Any, unit_seconds: float) -> dt.timedelta:
    if value in (None, ""):
        return dt.timedelta(0)
    raw = str(value).strip()
    if re.fullmatch(r"-?\d+(\.\d+)?", raw):
        return dt.timedelta(seconds=float(raw) * unit_seconds)
    if ":" in raw:
        sign = -1 if raw.startswith("-") else 1
        parts = [float(p) for p in raw.lstrip("-").split(":")]
        weights = [86400, 3600, 60, 1][-len(parts):]
        if len(parts) > 4:
            raise ValueError(f"无法解析时长: {raw}")
        return dt.timedelta(seconds=sign * sum(p * w for p, w in zip(parts, weights)))
    match = DURATION_RE.match(raw)
    if not match or not any(match.group(g) for g in "dhms"):
        raise ValueError(f"无法解析时长: {raw}")
    secs = sum(float(match.group(g) or 0) * w for g, w in (("d", 86400), ("h", 3600), ("m", 60), ("s", 1)))
    return dt.timedelta(seconds=-secs if match.group("sign") else secs)


def fmt_hours(delta: dt.timedelta | None) -> str:
    if delta is None:
        return ""
    hours = delta.total_seconds() / 3600
    return f"{fmt_num(hours)}h"


def fmt_num(value: float) -> str:
    return str(int(value)) if float(value).is_integer() else f"{value:.3f}".rstrip("0")


def has_cjk(value: str) -> bool:
    return any("\u4e00" <= ch <= "\u9fff" for ch in value)


# --------------------------------------------------------------------------- 工程定位


def is_project_root(path: Path) -> bool:
    return (path / JSON_REL).is_dir() or (path / "Shared" / "Data" / "Tables").is_dir()


def walk_up(start: Path) -> Path | None:
    for candidate in [start, *start.parents]:
        if is_project_root(candidate):
            return candidate
    return None


def python_test_config() -> dict | None:
    path = PYTHON_TEST_ROOT / PYTHON_TEST_CONFIG
    if not path.is_file():
        return None
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError):
        return None


def resolve_project_root(explicit: str | None, script_dir: Path) -> tuple[Path | None, str]:
    if explicit:
        root = Path(explicit).expanduser().resolve()
        return (root, "--project-root") if is_project_root(root) else (None, f"--project-root 不是 Beyond 工程: {root}")
    found = walk_up(script_dir)
    if found:
        return found, "skill 目录"
    found = walk_up(Path.cwd().resolve())
    if found:
        return found, "当前工作目录"
    config = python_test_config()
    if config:
        for key in ("project_root", "projectRoot", "beyond_root", "beyondRoot", "root"):
            value = config.get(key)
            if value and is_project_root(Path(value)):
                return Path(value), f"python-test {PYTHON_TEST_CONFIG}"
    return None, "未找到 Beyond 工程根（需含 Data/TableCfg/JSON 或 Shared/Data/Tables）"


def version_key(name: str) -> tuple[int, ...]:
    return tuple(int(n) for n in re.findall(r"\d+", name))


def resolve_json_dir(root: Path, version: str | None) -> tuple[str | None, Path | None]:
    base = root / JSON_REL
    if not base.is_dir():
        return version, None
    candidates = []
    for child in base.iterdir():
        match = JSON_DIR_RE.match(child.name)
        if child.is_dir() and match:
            if version and match.group(1).lower() != version.lower():
                continue
            candidates.append((version_key(child.name), match.group(1), child))
    if not candidates:
        return version, None
    candidates.sort()
    _, ver, path = candidates[-1]
    return version or ver, path


def resolve_output_dir(explicit: str | None, root: Path, script_dir: Path) -> Path:
    if explicit:
        return Path(explicit).expanduser().resolve()
    try:
        script_dir.resolve().relative_to(root.resolve())
        return root / REPORT_REL
    except ValueError:
        pass
    if python_test_config() is not None:
        return PYTHON_TEST_ROOT / "output" / "活动" / "reports"
    return root / REPORT_REL


# --------------------------------------------------------------------------- 读表


def extract_records(obj: Any) -> list[dict]:
    if isinstance(obj, list):
        return [x for x in obj if isinstance(x, dict)]
    if not isinstance(obj, dict):
        return []
    values = list(obj.values())
    single = values[0] if len(values) == 1 else None
    if isinstance(single, list) or (isinstance(single, dict) and "id" not in {str(k).lower() for k in single}):
        inner = extract_records(single)
        if inner:
            return inner
    if values and all(isinstance(v, dict) for v in values):
        return [{"_key": k, **v} for k, v in obj.items()]
    for value in values:
        if isinstance(value, (list, dict)):
            inner = extract_records(value)
            if inner:
                return inner
    return []


def find_table_files(json_dir: Path, names: list[str]) -> list[Path]:
    wanted = [n.lower() for n in names]
    hits = []
    for path in sorted(json_dir.rglob("*.json")):
        stem = path.stem.lower()
        if any(stem == n or stem.startswith(n + "_") or stem.startswith(n + ".") for n in wanted):
            hits.append(path)
    return hits


def load_table(json_dir: Path | None, names: list[str], id_aliases: list[str]) -> dict[str, dict]:
    if json_dir is None:
        return {}
    table: dict[str, dict] = {}
    for path in find_table_files(json_dir, names):
        try:
            data = json.loads(path.read_text(encoding="utf-8-sig"))
        except (OSError, json.JSONDecodeError):
            continue
        for rec in extract_records(data):
            key = text(pick(rec, id_aliases)) or text(rec.get("_key"))
            if key:
                table[key] = rec
    return table


def find_csv(root: Path) -> Path | None:
    base = root / CSV_CACHE_REL
    if not base.is_dir():
        return None
    hits = sorted(base.rglob(CSV_NAME))
    return hits[0] if hits else None


def read_csv_rows(path: Path, fields: dict[str, list[str]]) -> list[tuple[int, dict]]:
    raw = path.read_bytes()
    try:
        content = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        content = raw.decode("gb18030", errors="replace")
    rows = list(csv.reader(content.splitlines()))
    act_names = {a.lower() for a in fields["activity"]}
    type_names = {a.lower() for a in fields["push_type"]}
    header_idx = None
    for idx, row in enumerate(rows[:15]):
        cells = {c.strip().lower() for c in row}
        if cells & act_names and cells & type_names:
            header_idx = idx
            break
    if header_idx is None:
        raise ValueError(f"CSV 里找不到表头（需要 {fields['activity']} 和 {fields['push_type']} 之一）: {path}")
    header = [c.strip() for c in rows[header_idx]]
    out = []
    for idx, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        if not any(c.strip() for c in row):
            continue
        rec = {header[i]: row[i].strip() for i in range(min(len(header), len(row))) if header[i]}
        act = text(pick(rec, fields["activity"]))
        if not act or act.startswith("#") or act.lower() in META_TOKENS or has_cjk(act):
            continue
        out.append((idx, rec))
    return out


# --------------------------------------------------------------------------- 模型


@dataclass
class ServerTimes:
    open: dict[str, dt.datetime | None] = field(default_factory=dict)
    close: dict[str, dt.datetime | None] = field(default_factory=dict)


@dataclass
class Stage:
    name: str
    source: str
    times: ServerTimes


@dataclass
class Activity:
    id: str
    name: str
    times: ServerTimes
    version_start: str
    version_end: str
    stages: list[Stage]
    raw: dict


@dataclass
class Push:
    row: int
    id: str
    activity_id: str
    push_type: str
    kind: str | None
    sort: str
    offset: dt.timedelta | None
    end_raw: str
    version_start: str
    version_end: str
    activate: dict[str, dt.datetime | None] = field(default_factory=dict)
    until: dict[str, dt.datetime | None] = field(default_factory=dict)
    until_source: str = ""
    stage_ref: str = ""
    checked: bool = False

    @property
    def carrier(self) -> str | None:
        return self.kind.split("_")[0] if self.kind else None

    @property
    def phase(self) -> str | None:
        return self.kind.split("_")[1] if self.kind else None


@dataclass
class Issue:
    rule: str
    level: str
    activity_id: str
    push_ids: list[str]
    rows: list[int]
    message: str
    fix: str = ""


class Auditor:
    def __init__(self, schema: dict, json_dir: Path | None):
        self.schema = schema
        self.servers: dict[str, dict] = schema["servers"]
        self.tz = {k: dt.timezone(dt.timedelta(hours=v["utc_offset_hours"])) for k, v in self.servers.items()}
        self.primary = next(iter(self.servers))
        self.issues: list[Issue] = []
        tables = schema["tables"]
        self.time_ranges = load_table(json_dir, tables["time_range"], schema["time_range_fields"]["id"])
        self.activity_raw = load_table(json_dir, tables["activity"], schema["activity_fields"]["id"])
        self.cms = load_table(json_dir, tables["cms"], schema["stage_fields"]["id"])
        self.entrances = load_table(json_dir, tables["game_entrance"], schema["stage_fields"]["id"])
        self.activities: dict[str, Activity] = {}

    # ---- 时间解析

    def server_value(self, rec: dict, aliases: list[str], server: str) -> Any:
        suffixed = [a + s for a in aliases for s in self.servers[server].get("suffixes", [])]
        value = pick(rec, suffixed)
        return value if value is not None else pick(rec, aliases)

    def times_from(self, rec: dict, start_aliases: list[str], end_aliases: list[str]) -> ServerTimes:
        times = ServerTimes()
        for server in self.servers:
            times.open[server] = parse_time(self.server_value(rec, start_aliases, server), self.tz[server])
            times.close[server] = parse_time(self.server_value(rec, end_aliases, server), self.tz[server])
        return times

    def range_times(self, time_id: str) -> ServerTimes | None:
        rec = self.time_ranges.get(time_id)
        if rec is None:
            return None
        f = self.schema["time_range_fields"]
        times = self.times_from(rec, f["start"], f["end"])
        if all(v is None for v in times.open.values()):
            for value in rec.values():
                if isinstance(value, list) and value and isinstance(value[0], dict):
                    first = self.times_from(value[0], f["start"], f["end"])
                    last = self.times_from(value[-1], f["start"], f["end"])
                    return ServerTimes(open=first.open, close=last.close)
        return times

    def shifted(self, base: dict[str, dt.datetime | None], delta: dt.timedelta) -> dict[str, dt.datetime | None]:
        return {k: (v + delta if v else None) for k, v in base.items()}

    # ---- 活动与阶段

    def activity(self, act_id: str) -> Activity | None:
        if act_id in self.activities:
            return self.activities[act_id]
        rec = self.activity_raw.get(act_id)
        if rec is None:
            return None
        f = self.schema["activity_fields"]
        time_id = text(pick(rec, f["time_id"]))
        times = self.range_times(time_id) if time_id else None
        if times is None:
            times = self.times_from(rec, f["open"], f["close"])
        act = Activity(
            id=act_id,
            name=text(pick(rec, f["name"])),
            times=times,
            version_start=text(pick(rec, f["version_start"])),
            version_end=text(pick(rec, f["version_end"])),
            stages=self.collect_stages(act_id, time_id, times),
            raw=rec,
        )
        self.activities[act_id] = act
        return act

    def collect_stages(self, act_id: str, time_id: str, act_times: ServerTimes) -> list[Stage]:
        stages: list[Stage] = []
        unit = self.schema["offset_unit_seconds"]
        if time_id:
            pattern = re.compile(rf"^{re.escape(time_id)}_season_\d+_week_\d+", re.IGNORECASE)
            for rid in sorted(self.time_ranges, key=version_key):
                if pattern.match(rid):
                    times = self.range_times(rid)
                    if times:
                        stages.append(Stage(rid, "TimeRange season/week", times))
        f = self.schema["stage_fields"]
        for source, table in (("CMS", self.cms), ("GameEntrance", self.entrances)):
            for sid, rec in table.items():
                if text(pick(rec, f["activity"])) != act_id:
                    continue
                stage_time_id = text(pick(rec, f["time_id"])) if source == "CMS" else ""
                times = self.range_times(stage_time_id) if stage_time_id else None
                if times is None:
                    offset_raw = pick(rec, f["time_offset"])
                    if offset_raw is None:
                        continue
                    try:
                        offset = parse_duration(offset_raw, unit)
                    except ValueError:
                        continue
                    close_raw = pick(rec, f["close_offset"])
                    close = self.shifted(act_times.open, parse_duration(close_raw, unit)) if close_raw is not None else {}
                    times = ServerTimes(open=self.shifted(act_times.open, offset), close=close)
                stages.append(Stage(f"{source}:{sid}", source, times))
        seen: set = set()
        unique = []
        for stage in sorted(stages, key=lambda s: s.times.open.get(self.primary) or dt.datetime.max.replace(tzinfo=UTC)):
            key = (stage.times.open.get(self.primary), stage.times.close.get(self.primary))
            if key not in seen:
                seen.add(key)
                unique.append(stage)
        return unique

    # ---- push 行

    def classify(self, push_type: str, carrier: str) -> str | None:
        mapping = self.schema.get("push_type_map", {})
        if push_type in mapping:
            return mapping[push_type]
        blob = f"{carrier} {push_type}".lower()
        kind_carrier = "tab" if ("tab" in blob or "页签" in blob) else "bubble" if ("bubble" in blob or "气泡" in blob) else None
        if "update" in blob or "更新" in blob:
            phase = "update"
        elif any(k in blob for k in ("end", "临期", "expire", "close", "结束")):
            phase = "end"
        elif any(k in blob for k in ("open", "开放", "start")):
            phase = "open"
        else:
            phase = None
        if kind_carrier is None:
            kind_carrier = "bubble"
        return f"{kind_carrier}_{phase}" if phase else None

    def build_push(self, row: int, rec: dict) -> Push:
        f = self.schema["csv_fields"]
        push_type = text(pick(rec, f["push_type"]))
        offset_raw = pick(rec, f["offset"])
        push = Push(
            row=row,
            id=text(pick(rec, f["id"])) or f"row{row}",
            activity_id=text(pick(rec, f["activity"])),
            push_type=push_type,
            kind=self.classify(push_type, text(pick(rec, f["carrier"]))),
            sort=text(pick(rec, f["sort"])),
            offset=None,
            end_raw=text(pick(rec, f["end_time"])),
            version_start=text(pick(rec, f["version_start"])),
            version_end=text(pick(rec, f["version_end"])),
        )
        try:
            push.offset = parse_duration(offset_raw, self.schema["offset_unit_seconds"])
        except ValueError as exc:
            self.add("parse", "error", push, str(exc), "offset 填秒数，或 3d / 72h / d:hh:mm:ss")
        act = self.activity(push.activity_id)
        if act is None:
            self.add("activity_missing", "error", push, f"活动 {push.activity_id} 不在 Activity 表", "检查 activityId 或 JSON 是否已导出")
            return push
        if push.offset is not None:
            push.activate = self.shifted(act.times.open, push.offset)
        if push.end_raw:
            ranged = self.range_times(push.end_raw)
            if ranged:
                push.until, push.until_source = ranged.close, f"activityEndTime→{push.end_raw}"
            else:
                try:
                    push.until = {s: parse_time(push.end_raw, self.tz[s]) for s in self.servers}
                    push.until_source = "activityEndTime"
                except ValueError as exc:
                    self.add("parse", "error", push, str(exc), "activityEndTime 填时间或 TimeRange id")
        if not push.until:
            push.until, push.until_source = dict(act.times.close), "活动 closeTime"
        if push.activate.get(self.primary) is None or push.until.get(self.primary) is None:
            self.add("time_missing", "error", push, "算不出国服可弹 / 可弹至", "检查活动 timeId、TimeRange 与 offset")
        if push.kind is None:
            self.add("unknown_type", "warn", push, f"认不出 pushType={push.push_type!r}", f"在 {SCHEMA_FILE} 的 push_type_map 里登记")
        return push

    def add(self, rule: str, level: str, pushes: Push | list[Push], message: str, fix: str = "") -> None:
        items = pushes if isinstance(pushes, list) else [pushes]
        self.issues.append(
            Issue(rule, level, items[0].activity_id, [p.id for p in items], [p.row for p in items], message, fix)
        )

    # ---- 规则

    def fmt(self, value: dt.datetime | None, server: str | None = None) -> str:
        if value is None:
            return ""
        return value.astimezone(self.tz[server or self.primary]).strftime("%Y-%m-%d %H:%M")

    def run_rules(self, pushes: list[Push]) -> None:
        by_activity: dict[str, list[Push]] = {}
        for push in pushes:
            by_activity.setdefault(push.activity_id, []).append(push)
        for act_id, group in by_activity.items():
            if not any(p.checked for p in group):
                continue
            act = self.activities.get(act_id)
            self.rule_sort(group)
            self.rule_pair(group)
            for push in group:
                if not push.checked or act is None:
                    continue
                self.rule_times(push, act)

    def rule_sort(self, group: list[Push]) -> None:
        by_type: dict[str, list[Push]] = {}
        for push in group:
            if push.sort:
                by_type.setdefault(push.push_type, []).append(push)
        for push_type, items in by_type.items():
            seen: dict[str, list[Push]] = {}
            for push in items:
                seen.setdefault(push.sort, []).append(push)
            for sort, dup in seen.items():
                if len(dup) > 1 and any(p.checked for p in dup):
                    self.add("sort_dup", "error", dup, f"pushType={push_type} 的 bubbleSortId={sort} 重复", "同活动同 pushType 的 sort 改成互不相同")
        for carrier in ("bubble", "tab"):
            ends = [p for p in group if p.carrier == carrier and p.phase == "end" and p.sort]
            others = [p for p in group if p.carrier == carrier and p.phase in ("open", "update") and p.sort]
            for end in ends:
                for other in others:
                    try:
                        bad = float(end.sort) >= float(other.sort)
                    except ValueError:
                        continue
                    if bad and (end.checked or other.checked):
                        self.add(
                            "sort_order", "error", [end, other],
                            f"{KIND_LABEL[end.kind]} sort={end.sort} 不小于 {KIND_LABEL[other.kind]} sort={other.sort}",
                            "临期/End 的 bubbleSortId 要比已开放/Update 小",
                        )

    def rule_pair(self, group: list[Push]) -> None:
        kinds = {p.kind for p in group}
        for left, right in (("bubble_end", "tab_end"), ("bubble_update", "tab_update")):
            for have, miss in ((left, right), (right, left)):
                items = [p for p in group if p.kind == have]
                if have in kinds and miss not in kinds and any(p.checked for p in items):
                    self.add("pair", "warn", items, f"有 {KIND_LABEL[have]} 但没有 {KIND_LABEL[miss]}", f"确认是否漏配 {KIND_LABEL[miss]}")

    def rule_times(self, push: Push, act: Activity) -> None:
        primary = self.primary
        activate, until = push.activate.get(primary), push.until.get(primary)
        if activate is None or until is None:
            return
        late = [s for s in self.servers if push.activate.get(s) and push.until.get(s) and push.activate[s] > push.until[s]]
        if late:
            labels = "/".join(self.servers[s]["label"] for s in late)
            self.add("activate_after_end", "error", push, f"{labels} 激活晚于结束点", "缩短 offset 或延后结束点")

        apac = "APAC" if "APAC" in self.servers else None
        if apac and apac != primary:
            if push.activate.get(apac) != activate or push.until.get(apac) != until:
                self.add(
                    "tz_apac", "error", push,
                    f"亚太 {self.fmt(push.activate.get(apac), primary)}~{self.fmt(push.until.get(apac), primary)} "
                    f"≠ 国服 {self.fmt(activate)}~{self.fmt(until)}（按国服时区显示）",
                    "亚太时间要与国服绝对时刻一致",
                )

        if push.phase == "end":
            lead = dt.timedelta(hours=self.schema["end_lead_hours"])
            actual = until - activate
            if actual != lead:
                fix = ""
                open_time = act.times.open.get(primary)
                if open_time is not None:
                    want = (until - lead - open_time).total_seconds() / self.schema["offset_unit_seconds"]
                    fix = f"offset 应为 {fmt_num(want)}（激活 {self.fmt(until - lead)}）"
                self.add("end_minus_3d", "error", push, f"激活到结束 {fmt_hours(actual)}，应为 {fmt_hours(lead)}", fix)
            remain = {s: push.until[s] - push.activate[s] for s in self.servers if push.until.get(s) and push.activate.get(s)}
            if len(set(remain.values())) > 1:
                detail = "，".join(f"{self.servers[s]['label']} {fmt_hours(v)}" for s, v in remain.items())
                self.add("tz_end", "error", push, f"三服临期剩余时长不一致：{detail}", "三服结束前剩余时长要一致")
            closes = [st.times.close.get(primary) for st in act.stages if st.times.close.get(primary)]
            if closes:
                candidates = closes + [act.times.close.get(primary)]
                match = next((st for st in act.stages if st.times.close.get(primary) == until), None)
                if until not in candidates:
                    nearest = min(closes, key=lambda c: abs((c - until).total_seconds()))
                    self.add("end_vs_stage", "error", push, f"可弹至 {self.fmt(until)} 不等于任何阶段 closeTime", f"最近阶段结束 {self.fmt(nearest)}")
                push.stage_ref = match.name if match else "活动 closeTime"
            else:
                push.stage_ref = push.until_source

        if push.phase == "update":
            act_open = act.times.open.get(primary)
            opens = [st for st in act.stages if st.times.open.get(primary) and (act_open is None or st.times.open[primary] > act_open)]
            if not opens:
                self.add("update_vs_stage", "warn", push, "没找到后续阶段（CMS / season_week / GameEntrance），无法核 Update", "确认阶段来源或手工核对")
            else:
                match = next((st for st in opens if st.times.open[primary] == activate), None)
                if match is None:
                    nearest = min(opens, key=lambda st: abs((st.times.open[primary] - activate).total_seconds()))
                    want = ""
                    if act_open is not None:
                        want = f"，offset 应为 {fmt_num((nearest.times.open[primary] - act_open).total_seconds() / self.schema['offset_unit_seconds'])}"
                    self.add(
                        "update_vs_stage", "error", push,
                        f"Update 激活 {self.fmt(activate)} 不等于任何后续阶段开放时刻",
                        f"最近阶段 {nearest.name} 开放 {self.fmt(nearest.times.open[primary])}{want}",
                    )
                else:
                    push.stage_ref = match.name
            servers = [s for s in self.servers if push.activate.get(s)]
            wall = {push.activate[s].astimezone(self.tz[s]).replace(tzinfo=None) for s in servers}
            since_open = {push.activate[s] - act.times.open[s] for s in servers if act.times.open.get(s)}
            if len(wall) > 1 and len(since_open) > 1:
                detail = "，".join(f"{self.servers[s]['label']} {self.fmt(push.activate[s], s)}" for s in servers)
                self.add("tz_update", "error", push, f"三服 Update 既不墙钟对齐也不按开服后时长对齐：{detail}", "统一用墙钟或统一用开服后 offset")


# --------------------------------------------------------------------------- 输出


def version_window(auditor: Auditor, version: str | None) -> ServerTimes | None:
    if not version:
        return None
    return auditor.range_times(f"time_version_{version}")


def overview_rows(auditor: Auditor, pushes: list[Push], window: ServerTimes | None, all_mode: bool) -> list[dict]:
    primary = auditor.primary
    start = window.open.get(primary) if window else None
    end = window.close.get(primary) if window else None
    chosen: list[Push] = []
    for push in pushes:
        act = auditor.activities.get(push.activity_id)
        if act is None:
            continue
        if all_mode or start is None:
            chosen.append(push)
            continue
        act_open = act.times.open.get(primary)
        activate, until = push.activate.get(primary), push.until.get(primary)
        if act_open is not None and act_open >= start:
            if until is None or until > start:
                chosen.append(push)
        elif activate is not None and activate >= start and (end is None or activate < end):
            chosen.append(push)
    far = dt.datetime.max.replace(tzinfo=UTC)
    chosen.sort(
        key=lambda p: (
            auditor.activities[p.activity_id].times.open.get(primary) or far,
            p.activity_id,
            p.activate.get(primary) or far,
        )
    )
    other = next((s for s in auditor.servers if s != primary and s != "APAC"), None)
    rows = []
    for push in chosen:
        act = auditor.activities[push.activity_id]
        vs = act.version_start or push.version_start
        ve = act.version_end or push.version_end
        row = {
            "活动版本": f"{vs} ~ {ve}" if ve else f"{vs} ~ 常驻",
            "活动名": act.name,
            "活动id": act.id,
            "pushId": push.id,
            "类型": KIND_LABEL.get(push.kind or "", push.push_type),
            "sort": push.sort,
            "参照阶段": push.stage_ref or ("活动开启" if push.phase == "open" else ""),
            "国服可弹": auditor.fmt(push.activate.get(primary)),
            "国服可弹至": auditor.fmt(push.until.get(primary)),
        }
        if other:
            label = auditor.servers[other]["label"]
            row[f"{label}触发"] = auditor.fmt(push.activate.get(other), other)
            row[f"{label}结束"] = auditor.fmt(push.until.get(other), other)
        rows.append(row)
    return rows


def push_dict(auditor: Auditor, push: Push, levels: dict[int, str]) -> dict:
    act = auditor.activities.get(push.activity_id)
    return {
        "row": push.row,
        "pushId": push.id,
        "activityId": push.activity_id,
        "activityName": act.name if act else "",
        "pushType": push.push_type,
        "kind": push.kind,
        "kindLabel": KIND_LABEL.get(push.kind or "", ""),
        "sort": push.sort,
        "versionStart": push.version_start,
        "offsetSeconds": push.offset.total_seconds() if push.offset is not None else None,
        "activate": {s: auditor.fmt(v, s) for s, v in push.activate.items()},
        "until": {s: auditor.fmt(v, s) for s, v in push.until.items()},
        "untilSource": push.until_source,
        "remain": fmt_hours(push.until.get(auditor.primary) - push.activate.get(auditor.primary))
        if push.until.get(auditor.primary) and push.activate.get(auditor.primary)
        else "",
        "stageRef": push.stage_ref,
        "result": levels.get(push.row, "ok"),
    }


def md_escape(value: Any) -> str:
    return text(value).replace("|", "\\|").replace("\n", " ")


def md_table(headers: list[str], rows: list[list[Any]]) -> list[str]:
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
    lines += ["| " + " | ".join(md_escape(c) for c in row) + " |" for row in rows]
    return lines


def write_markdown(path: Path, meta: dict, items: list[dict], issues: list[Issue], overview: list[dict]) -> None:
    lines = [f"# {meta['title']}", ""]
    lines += [f"- {k}: `{v}`" for k, v in meta["info"].items()]
    lines += ["", "## 摘要", ""]
    lines += md_table(["级别", "数量"], [["error", meta["errors"]], ["warn", meta["warns"]], ["核对行数", len(items)]])
    lines += ["", "### 按规则", ""]
    lines += md_table(["规则", "error", "warn"], [[r, c["error"], c["warn"]] for r, c in meta["by_rule"].items()]) if meta["by_rule"] else ["无问题。"]
    lines += ["", "## 问题", ""]
    if issues:
        lines += md_table(
            ["级别", "规则", "活动", "pushId", "行", "说明", "建议"],
            [[i.level, i.rule, i.activity_id, ",".join(i.push_ids), ",".join(map(str, i.rows)), i.message, i.fix] for i in issues],
        )
    else:
        lines.append("无。")
    lines += ["", "## 核对明细", ""]
    servers = list(meta["servers"])
    headers = ["行", "活动id", "活动名", "pushId", "类型", "sort"]
    headers += [f"{meta['servers'][s]}可弹" for s in servers] + [f"{meta['servers'][servers[0]]}可弹至", "剩余", "参照阶段", "结果"]
    lines += md_table(
        headers,
        [
            [i["row"], i["activityId"], i["activityName"], i["pushId"], i["kindLabel"] or i["pushType"], i["sort"]]
            + [i["activate"].get(s, "") for s in servers]
            + [i["until"].get(servers[0], ""), i["remain"], i["stageRef"], i["result"]]
            for i in items
        ],
    )
    if overview:
        lines += ["", "## 一览（全量）", ""]
        lines += md_table(list(overview[0]), [list(r.values()) for r in overview])
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_xlsx(path: Path, meta: dict, items: list[dict], issues: list[Issue], overview: list[dict]) -> str | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return "未安装 openpyxl（pip install openpyxl），跳过 xlsx"
    wb = Workbook()
    bold = Font(bold=True)
    fills = {"error": PatternFill("solid", fgColor="F8CBAD"), "warn": PatternFill("solid", fgColor="FFE699")}

    def sheet(title: str, headers: list[str], rows: list[list[Any]], first: bool = False):
        ws = wb.active if first else wb.create_sheet(title)
        ws.title = title
        ws.append(headers)
        for cell in ws[1]:
            cell.font = bold
        for row in rows:
            ws.append(row)
        for col in ws.columns:
            width = max(len(text(c.value)) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width * 1.6), 60)
        ws.freeze_panes = "A2"
        return ws

    headers = list(overview[0]) if overview else ["活动版本", "活动名", "活动id", "pushId", "类型"]
    ws = sheet("全量", headers, [list(r.values()) for r in overview], first=True)
    first = 0
    for idx in range(1, len(overview) + 1):
        if idx < len(overview) and overview[idx]["活动id"] == overview[first]["活动id"]:
            continue
        if idx - 1 > first:
            for col in (1, 2, 3):
                ws.merge_cells(start_row=first + 2, start_column=col, end_row=idx + 1, end_column=col)
                ws.cell(first + 2, col).alignment = Alignment(vertical="center")
        first = idx

    def kind_rows(phase: str) -> list[list[Any]]:
        primary = meta["primary"]
        return [
            [i["activityId"], i["activityName"], i["pushId"], i["kindLabel"], i["activate"].get(primary, ""),
             i["until"].get(primary, ""), i["remain"], i["stageRef"], i["result"]]
            for i in items
            if (i["kind"] or "").endswith(phase)
        ]

    cols = ["活动id", "活动名", "pushId", "类型", "国服可弹", "国服可弹至", "剩余", "参照阶段", "结果"]
    for title, phase in (("Update对齐", "_update"), ("结束对齐", "_end")):
        ws = sheet(title, cols, kind_rows(phase))
        for row in ws.iter_rows(min_row=2):
            fill = fills.get(text(row[-1].value))
            if fill:
                for cell in row:
                    cell.fill = fill
    ws = sheet(
        "问题",
        ["级别", "规则", "活动", "pushId", "行", "说明", "建议"],
        [[i.level, i.rule, i.activity_id, ",".join(i.push_ids), ",".join(map(str, i.rows)), i.message, i.fix] for i in issues],
    )
    for row in ws.iter_rows(min_row=2):
        fill = fills.get(text(row[0].value))
        if fill:
            for cell in row:
                cell.fill = fill
    summary = [[k, v] for k, v in meta["info"].items()] + [["error", meta["errors"]], ["warn", meta["warns"]], ["核对行数", len(items)]]
    summary += [[f"规则 {r}", f"error {c['error']} / warn {c['warn']}"] for r, c in meta["by_rule"].items()]
    sheet("摘要", ["项", "值"], summary)
    wb.save(path)
    return None


# --------------------------------------------------------------------------- 入口


def load_schema(script_dir: Path, explicit: str | None) -> tuple[dict, str]:
    schema = DEFAULT_SCHEMA
    path = Path(explicit) if explicit else script_dir / SCHEMA_FILE
    if path.is_file():
        schema = merge_schema(schema, json.loads(path.read_text(encoding="utf-8-sig")))
        return schema, str(path)
    if explicit:
        raise FileNotFoundError(f"--schema 不存在: {path}")
    return schema, "内置默认"


def parse_args(argv: list[str] | None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Beyond 活动 Push 气泡 / Tab 配置核对（只读）")
    parser.add_argument("--project-root", help="Beyond 工程根；默认从 skill 目录 / 当前目录向上找")
    parser.add_argument("--current-version", help="大版本，如 v1d6；默认从 Data/TableCfg/JSON/vXdYd99* 推出")
    parser.add_argument("--all", action="store_true", help="扫全表，不按 versionStart 过滤")
    parser.add_argument("--output-dir", help="报告目录，覆盖默认位置")
    parser.add_argument("--schema", help=f"字段映射 JSON，默认读脚本同目录 {SCHEMA_FILE}")
    parser.add_argument("--fail-on-error", action="store_true", help="有 error 时退出码为 1")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass
    args = parse_args(argv)
    script_dir = Path(__file__).resolve().parent
    try:
        schema, schema_source = load_schema(script_dir, args.schema)
    except (OSError, json.JSONDecodeError) as exc:
        print(f"error=schema 读取失败: {exc}")
        return 2

    root, root_source = resolve_project_root(args.project_root, script_dir)
    if root is None:
        print(f"error={root_source}")
        return 2
    print(f"project={root}  (来源: {root_source})")

    version, json_dir = resolve_json_dir(root, args.current_version)
    if json_dir is None:
        print(f"error=找不到 {root / JSON_REL}/{version or 'vXdY'}d99* 导出目录")
        return 2
    if not args.all and not version:
        print("error=推不出大版本，请传 --current-version 或 --all")
        return 2
    csv_path = find_csv(root)
    if csv_path is None:
        print(f"error=找不到 {root / CSV_CACHE_REL}/**/{CSV_NAME}")
        return 2
    print(f"version={version}  json={json_dir}")
    print(f"csv={csv_path}")

    auditor = Auditor(schema, json_dir)
    if not auditor.activity_raw:
        print(f"error=JSON 目录里没读到 Activity 表（{schema['tables']['activity']}）")
        return 2
    try:
        raw_rows = read_csv_rows(csv_path, schema["csv_fields"])
    except ValueError as exc:
        print(f"error={exc}")
        return 2

    pushes = [auditor.build_push(row, rec) for row, rec in raw_rows]
    version_re = re.compile(rf"^{re.escape(version)}(d\d+)?$", re.IGNORECASE) if version else None
    for push in pushes:
        push.checked = args.all or bool(
            CURRENT_RE.match(push.version_start) or (version_re and version_re.match(push.version_start))
        )
    checked_rows = {p.row for p in pushes if p.checked}
    auditor.issues = [i for i in auditor.issues if any(r in checked_rows for r in i.rows)]
    auditor.run_rules(pushes)

    issues = sorted(auditor.issues, key=lambda i: (i.level != "error", i.activity_id, min(i.rows), i.rule))
    levels: dict[int, str] = {}
    for issue in issues:
        for row in issue.rows:
            if row in checked_rows and levels.get(row) != "error":
                levels[row] = issue.level
    checked = [p for p in pushes if p.checked]
    items = [push_dict(auditor, p, levels) for p in checked]
    overview = overview_rows(auditor, pushes, version_window(auditor, version), args.all)

    by_rule: dict[str, dict[str, int]] = {}
    for issue in issues:
        by_rule.setdefault(issue.rule, {"error": 0, "warn": 0})[issue.level] += 1
    errors = sum(1 for i in issues if i.level == "error")
    warns = sum(1 for i in issues if i.level == "warn")
    prefix = "all" if args.all else version
    today = dt.datetime.now(dt.timezone(dt.timedelta(hours=8))).strftime("%Y%m%d")
    stem = f"{prefix}_push气泡Tab配置核对_{today}"
    meta = {
        "title": f"{prefix} Push 气泡 / Tab 配置核对",
        "info": {
            "project": str(root),
            "version": version or "",
            "mode": "all" if args.all else "version",
            "json": str(json_dir),
            "csv": str(csv_path),
            "schema": schema_source,
            "生成时间(UTC+8)": dt.datetime.now(dt.timezone(dt.timedelta(hours=8))).strftime("%Y-%m-%d %H:%M"),
        },
        "errors": errors,
        "warns": warns,
        "by_rule": by_rule,
        "servers": {k: v["label"] for k, v in schema["servers"].items()},
        "primary": auditor.primary,
    }

    out_dir = resolve_output_dir(args.output_dir, root, script_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    md_path, json_path, xlsx_path = (out_dir / f"{stem}{ext}" for ext in (".md", ".json", ".xlsx"))
    write_markdown(md_path, meta, items, issues, overview)
    json_path.write_text(
        json.dumps(
            {**meta, "issues": [i.__dict__ for i in issues], "items": items, "overview": overview},
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    xlsx_note = write_xlsx(xlsx_path, meta, items, issues, overview)

    print(f"mode={'all' if args.all else 'version'}  rows={len(pushes)}  checked={len(checked)}  activities={len({p.activity_id for p in checked})}")
    print(f"summary: error={errors} warn={warns}")
    for rule, counts in by_rule.items():
        print(f"  {rule}: error={counts['error']} warn={counts['warn']}")
    for issue in [i for i in issues if i.level == "error"][:20]:
        fix = f" → {issue.fix}" if issue.fix else ""
        print(f"  [error] {issue.rule} 活动={issue.activity_id} push={','.join(issue.push_ids)} 行={','.join(map(str, issue.rows))}: {issue.message}{fix}")
    if errors > 20:
        print(f"  ... 其余 {errors - 20} 条 error 见报告")
    print(f"report_md={md_path}")
    print(f"report_json={json_path}")
    print(f"report_xlsx={xlsx_note or xlsx_path}")
    return 1 if args.fail_on_error and errors else 0


if __name__ == "__main__":
    sys.exit(main())
